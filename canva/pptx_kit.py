"""Utilidades para construir laminas 1920x1080 en PPTX para importarlas a Canva.

Todas las coordenadas se expresan en pixeles del lienzo de 1920x1080; el modulo
las convierte a EMU (1 px = 9525 EMU, 96 ppp). Los tamanos de letra tambien se
expresan en pixeles y se convierten a puntos (1 px = 0,75 pt).
"""

from __future__ import annotations

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt
from lxml import etree

W, H = 1920, 1080

FONT_HEAD = "Montserrat"
FONT_BODY = "Inter"
FONT_MONO = "Roboto Mono"


def px(v: float) -> Emu:
    return Emu(int(round(v * 9525)))


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#").upper())


def _set_alpha(color_parent, alpha: float) -> None:
    """Agrega <a:alpha> al srgbClr dentro de color_parent (fill o line)."""
    if alpha is None or alpha >= 1:
        return
    srgb = color_parent.find(qn("a:srgbClr"))
    if srgb is None:
        srgb = color_parent.find(".//" + qn("a:srgbClr"))
    if srgb is None:
        return
    for old in srgb.findall(qn("a:alpha")):
        srgb.remove(old)
    el = etree.SubElement(srgb, qn("a:alpha"))
    el.set("val", str(int(round(alpha * 100000))))


def new_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = px(W)
    prs.slide_height = px(H)
    return prs


def add_slide(prs: Presentation, bg: str = "#07172D", bg_image: str | None = None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(bg)
    if bg_image:
        pic = slide.shapes.add_picture(bg_image, 0, 0, px(W), px(H))
        pic.name = "Fondo"
    return slide


def notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def _style_fill(shape, fill: str | None, fill_alpha: float = 1.0) -> None:
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        if fill_alpha < 1:
            sp_pr = shape._element.spPr
            _set_alpha(sp_pr.find(qn("a:solidFill")), fill_alpha)


def _style_line(shape, line: str | None, width: float = 1.0, alpha: float = 1.0) -> None:
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(width * 0.75)
        if alpha < 1:
            ln = shape._element.spPr.find(qn("a:ln"))
            _set_alpha(ln.find(qn("a:solidFill")), alpha)


def _no_shadow(shape) -> None:
    sp_pr = shape._element.spPr
    if sp_pr.find(qn("a:effectLst")) is None:
        etree.SubElement(sp_pr, qn("a:effectLst"))


def rect(slide, x, y, w, h, fill=None, line=None, line_w=1.0, fill_alpha=1.0,
         line_alpha=1.0, radius: float | None = None, name: str | None = None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(kind, px(x), px(y), px(w), px(h))
    if radius:
        shp.adjustments[0] = min(0.5, radius / min(w, h))
    _style_fill(shp, fill, fill_alpha)
    _style_line(shp, line, line_w, line_alpha)
    _no_shadow(shp)
    shp.text_frame.text = ""
    if name:
        shp.name = name
    return shp


def ellipse(slide, cx, cy, r, fill=None, line=None, line_w=1.0, fill_alpha=1.0,
            line_alpha=1.0, name: str | None = None):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, px(cx - r), px(cy - r), px(2 * r), px(2 * r))
    _style_fill(shp, fill, fill_alpha)
    _style_line(shp, line, line_w, line_alpha)
    _no_shadow(shp)
    if name:
        shp.name = name
    return shp


def line(slide, x1, y1, x2, y2, color="#6D8FB2", width=1.0, alpha=1.0, arrow=False,
         dash: str | None = None):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, px(x1), px(y1), px(x2), px(y2))
    conn.line.color.rgb = rgb(color)
    conn.line.width = Pt(width * 0.75)
    ln = conn._element.spPr.find(qn("a:ln"))
    if alpha < 1:
        _set_alpha(ln.find(qn("a:solidFill")), alpha)
    if dash:
        d = etree.SubElement(ln, qn("a:prstDash"))
        d.set("val", dash)
    if arrow:
        tail = etree.SubElement(ln, qn("a:tailEnd"))
        tail.set("type", "triangle")
        tail.set("w", "med")
        tail.set("len", "med")
    return conn


def corners(slide, x, y, w, h, size=34, color="#FFD21F", width=3.0, which="tl,tr,bl,br"):
    """Esquineros en L (marcas de encuadre)."""
    parts = set(which.split(","))
    if "tl" in parts:
        line(slide, x, y, x + size, y, color, width)
        line(slide, x, y, x, y + size, color, width)
    if "tr" in parts:
        line(slide, x + w - size, y, x + w, y, color, width)
        line(slide, x + w, y, x + w, y + size, color, width)
    if "bl" in parts:
        line(slide, x, y + h - size, x, y + h, color, width)
        line(slide, x, y + h, x + size, y + h, color, width)
    if "br" in parts:
        line(slide, x + w, y + h - size, x + w, y + h, color, width)
        line(slide, x + w - size, y + h, x + w, y + h, color, width)


def text(slide, x, y, w, h, content, size=28, color="#F3F7FC", bold=False,
         font: str = FONT_BODY, align="left", anchor="top", line_spacing=1.15,
         italic=False, spacing: float | None = None, name: str | None = None,
         wrap=True):
    """Caja de texto. `content` puede ser str o lista de parrafos.

    Cada parrafo puede ser str o lista de runs; cada run es str o dict con
    claves text, color, bold, size, font, italic.
    """
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE,
                          "bottom": MSO_ANCHOR.BOTTOM}[anchor]
    paragraphs = content if isinstance(content, list) else [content]
    for i, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
                       "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}[align]
        p.line_spacing = line_spacing
        runs = para if isinstance(para, list) else [para]
        for run_spec in runs:
            spec = run_spec if isinstance(run_spec, dict) else {"text": run_spec}
            r = p.add_run()
            r.text = spec["text"]
            f = r.font
            f.size = Pt(spec.get("size", size) * 0.75)
            f.bold = spec.get("bold", bold)
            f.italic = spec.get("italic", italic)
            f.name = spec.get("font", font)
            f.color.rgb = rgb(spec.get("color", color))
            sp = spec.get("spacing", spacing)
            if sp:
                r._r.get_or_add_rPr().set("spc", str(int(sp * 100)))
    if name:
        tb.name = name
    return tb


def image(slide, path, x, y, w=None, h=None, name: str | None = None):
    pic = slide.shapes.add_picture(path, px(x), px(y), px(w) if w else None, px(h) if h else None)
    if name:
        pic.name = name
    return pic


def image_fit(slide, path, x, y, w, h, name: str | None = None):
    """Inserta la imagen completa (contain) centrada dentro de la caja."""
    from PIL import Image

    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    dw, dh = iw * scale, ih * scale
    return image(slide, path, x + (w - dw) / 2, y + (h - dh) / 2, dw, dh, name=name)
